"""
Evaluador de recuperación de información (RAG) con Claude como juez y validación humana.
Inspirado en eval-qna-rag-metrics (johntango/PromptFlowEval), sin Prompt Flow ni Azure.

EXCEL (tenis_eval.xlsx):
  Hoja "Fuentes"    : los PDF que son la fuente de verdad (archivo, titulo, autor, notas)
  Hoja "Preguntas"  : tus preguntas y lo que respondió ChatGPT (id, question, answer)
                      + columnas que llena el script y que tú validas

FLUJO:
  0. py eval_rag.py plantilla      -> crea tenis_eval.xlsx de ejemplo
  1. Llenas "Fuentes" con tus 3 PDF y "Preguntas" con question + answer
  2. py eval_rag.py sugerir         -> busca en TODOS los PDF y llena:
                                       fuente_sugerida, pagina, fragmento_sugerido, cita_verificada
                                       (y precarga "fuente" y "documents")
  3. TÚ REVISAS cada fila: corriges si hace falta y pones SI en "validado".
     Si la respuesta no está en ningún manual: "documents" vacío y SI.
  4. py eval_rag.py evaluar         -> evalúa solo filas validadas -> resultados.xlsx

Métricas del juez (1 a 5):
  retrieval_score : ¿el fragmento sirve para responder la pregunta?
  groundedness    : ¿lo que dice la respuesta está respaldado por el fragmento?
  relevance       : ¿la respuesta cubre lo que diría una referencia construida solo con el fragmento?

Requisitos:  py -m pip install anthropic python-dotenv openpyxl pypdf
Clave:       archivo .env con  ANTHROPIC_API_KEY=sk-ant-...   (opcional: JUDGE_MODEL=...)
Prueba:      agrega --dry-run a cualquier comando para no llamar a la API.
"""

import argparse
import difflib
import json
import math
import os
import random
import re
import sys
import time
import unicodedata
from collections import Counter
from pathlib import Path

try:
    from dotenv import load_dotenv
    load_dotenv()
except ImportError:
    pass

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
except ImportError:
    sys.exit("Falta el paquete: py -m pip install openpyxl")

DEFAULT_MODEL = os.getenv("JUDGE_MODEL", "claude-sonnet-5-5")
SH_FUENTES, SH_PREG = "Fuentes", "Preguntas"
COLS_FUENTES = ["archivo", "titulo", "autor", "notas"]
COLS_PREG = ["id", "question", "answer", "fuente_sugerida", "pagina", "fragmento_sugerido",
             "cita_verificada", "fuente", "documents", "validado"]
METRICS_ALL = ["retrieval_score", "groundedness", "relevance"]
NO_FUENTE = "(ninguna)"

# =========================================================================== prompts

SYSTEM_JUDGE = (
    "Eres un evaluador estricto e imparcial de sistemas de pregunta-respuesta basados en documentos. "
    "Calificas únicamente con la rúbrica que se te da. No usas tu conocimiento general para dar por "
    "buena una afirmación: si no está en los documentos, no está respaldada. "
    'Respondes SOLO con un objeto JSON: {"score": <entero 1-5>, "reason": "<una o dos frases en español>"}'
)

RUBRICS = {
    "retrieval_score": """Evalúa qué tan útiles son los DOCUMENTOS para responder la PREGUNTA.
1 = ningún documento es relevante para la pregunta (o no hay documentos).
2 = hay relación temática, pero no contienen la información necesaria.
3 = contienen parte de la información necesaria.
4 = contienen casi toda la información necesaria.
5 = uno o la combinación de varios documentos es ideal para responder la pregunta.

PREGUNTA:
{question}

DOCUMENTOS:
{documents}""",

    "groundedness": """Evalúa si las afirmaciones de la RESPUESTA están respaldadas por los DOCUMENTOS.
Aunque una afirmación sea verdadera en el mundo real, si no se puede verificar en los documentos
se considera NO respaldada. Si la respuesta dice correctamente que la información no está en los
documentos y efectivamente no está, califica 5.
1 = la respuesta es mayormente inventada o contradice los documentos.
2 = la mayoría de afirmaciones no están respaldadas.
3 = mezcla de afirmaciones respaldadas y no respaldadas.
4 = casi todo está respaldado; detalles menores añadidos.
5 = todas las afirmaciones están respaldadas por los documentos.

DOCUMENTOS:
{documents}

PREGUNTA:
{question}

RESPUESTA A EVALUAR:
{answer}""",

    "relevance": """Compara la RESPUESTA con la RESPUESTA DE REFERENCIA (construida solo con los documentos).
1 = la respuesta es completamente irrelevante frente a la referencia.
2 = toca el tema pero omite lo esencial o es incorrecta.
3 = cubre parte de lo esencial.
4 = cubre casi todo lo esencial.
5 = incluye toda la información necesaria para responder según la referencia.
Si la referencia indica que los documentos no permiten responder y la respuesta también lo reconoce,
califica 5; si en ese caso la respuesta inventa una contestación, califica 1.

PREGUNTA:
{question}

RESPUESTA DE REFERENCIA:
{reference}

RESPUESTA A EVALUAR:
{answer}""",
}

REFERENCE_PROMPT = """Responde la PREGUNTA usando EXCLUSIVAMENTE los DOCUMENTOS.
- Lee con cuidado: la respuesta puede estar en una regla, un inciso, un caso con su decisión
  (p. ej. "Caso: ¿se permite...? Decisión: Sí"), una nota o una excepción, aunque use palabras
  distintas a la pregunta (ej.: "suelo/tierra" = "terreno"; "el que saca" = "servidor").
- Puedes deducir lo que se sigue directamente del texto, pero no agregues datos que no estén.
- Solo si de verdad no hay nada que responda la pregunta, responde exactamente:
  "Los documentos no contienen información para responder esta pregunta."
Sé breve y preciso: una a tres oraciones.

DOCUMENTOS:
{documents}

PREGUNTA:
{question}"""

KEYWORDS_PROMPT = """Vas a buscar en manuales de deporte (reglamentos oficiales, manuales técnicos o de
entrenamiento; pueden estar en español o en inglés) el pasaje que responde esta pregunta.
La pregunta usa lenguaje coloquial, pero el manual usa vocabulario formal o técnico. Devuelve:
- palabras clave de la pregunta y sus sinónimos,
- cómo lo diría un reglamento o manual técnico (ej.: "el que saca" -> "servidor"; "suelo/tierra" -> "terreno", "pista";
  "pisar la línea" -> "falta de pie"),
- el nombre probable de la regla o sección,
- la traducción al inglés de los términos principales.
Máximo 25 términos, cada uno de 1 a 3 palabras.
Responde SOLO con JSON: {{"terminos": ["...", "..."]}}

PREGUNTA:
{question}"""

LOCATE_PROMPT = """Abajo hay fragmentos extraídos de varios manuales, cada uno con su archivo y página.
Identifica el o los pasajes que responden la PREGUNTA.

Reglas:
- Copia cada pasaje TEXTUALMENTE, tal como aparece (sin traducir, resumir ni corregir), de 1 a 4 oraciones.
- Usa un solo pasaje si basta. Usa hasta 3 solo si la respuesta está repartida entre manuales o páginas.
- La respuesta puede estar en una regla, pero también en casos, ejemplos, notas, excepciones o decisiones
  (p. ej. "Caso: ¿se permite...? Decisión: Sí"). Eso cuenta como respuesta aunque use otras palabras que la pregunta.
  En ese caso copia el caso completo con su decisión.
- Revisa TODOS los fragmentos antes de concluir que no hay respuesta.
- Si ningún fragmento responde la pregunta, devuelve "pasajes": [] . No inventes.

Responde SOLO con JSON:
{{"pasajes": [{{"archivo": "<nombre exacto>", "pagina": <número>, "fragmento": "<cita textual>"}}],
  "comentario": "<una frase: dónde está la respuesta o por qué no se encontró>"}}

PREGUNTA:
{question}

FRAGMENTOS:
{chunks}"""


# =========================================================================== cliente

class Claude:
    def __init__(self, model: str, dry_run: bool = False):
        self.model, self.dry_run, self.use_temperature = model, dry_run, True
        if dry_run:
            return
        try:
            import anthropic
        except ImportError:
            sys.exit("Falta el paquete: py -m pip install anthropic")
        if not os.getenv("ANTHROPIC_API_KEY"):
            sys.exit("No encuentro ANTHROPIC_API_KEY. Revisa el archivo .env (y py -m pip install python-dotenv).")
        self.anthropic = anthropic
        self.client = anthropic.Anthropic()

    def call(self, prompt: str, system: str = "", max_tokens: int = 500) -> str:
        kwargs = dict(model=self.model, max_tokens=max_tokens,
                      messages=[{"role": "user", "content": prompt}])
        if system:
            kwargs["system"] = system
        for attempt in range(5):
            try:
                if self.use_temperature:
                    kwargs["temperature"] = 0
                else:
                    kwargs.pop("temperature", None)
                resp = self.client.messages.create(**kwargs)
                return "".join(b.text for b in resp.content if getattr(b, "type", "") == "text")
            except self.anthropic.BadRequestError as e:
                if self.use_temperature and "temperature" in str(e).lower():
                    self.use_temperature = False      # algunos modelos no aceptan temperature
                    continue
                raise
            except (self.anthropic.RateLimitError, self.anthropic.APIConnectionError,
                    self.anthropic.InternalServerError):
                time.sleep(3 * 2 ** attempt)
        raise RuntimeError("La API no respondió tras varios intentos.")


def parse_json(text: str):
    m = re.search(r"\{.*\}", text or "", re.DOTALL)
    if not m:
        return None
    try:
        return json.loads(m.group(0))
    except json.JSONDecodeError:
        return None


def parse_score(text: str):
    """Extrae score y razón; tolera JSON incompleto (respuesta cortada) o texto alrededor."""
    text = text or ""
    obj = parse_json(text)
    if obj:
        try:
            s = int(obj.get("score"))
            if 1 <= s <= 5:
                return s, str(obj.get("reason", "")).strip()
        except (TypeError, ValueError):
            pass
    m = re.search(r'"score"\s*:\s*"?([1-5])', text)
    if m:
        r = re.search(r'"reason"\s*:\s*"(.*)', text, re.DOTALL)
        reason = re.sub(r'"\s*\}?\s*$', "", r.group(1)).strip() if r else ""
        return int(m.group(1)), reason + (" [razón incompleta]" if r and not text.rstrip().endswith("}") else "")
    return None, text.strip()[:300]


# =========================================================================== texto / PDF

STOP = set("""a al algo ante con contra cual cuales cuando de del desde donde e el ella ellos en entre
era es esa ese eso esta este esto estos fue ha hay la las le les lo los mas me mi muy no nos o para
pero por que se segun ser si sin sobre su sus tambien te tiene un una uno unos y ya como cuanto
cuantos cuantas que dice manual documento debe deben segun the of and to in is a an for on with
be are as by it or at this that from""".split())


def norm(s: str) -> str:
    s = unicodedata.normalize("NFKD", s or "")
    s = "".join(c for c in s if not unicodedata.combining(c)).lower()
    s = s.translate(str.maketrans({"“": '"', "”": '"', "‘": "'", "’": "'", "–": "-", "—": "-"}))
    return re.sub(r"\s+", " ", s).strip()


def tokens(s: str):
    return [t for t in re.findall(r"[a-z0-9]+", norm(s)) if len(t) > 1 and t not in STOP]


def load_pdf(path: Path):
    try:
        from pypdf import PdfReader
    except ImportError:
        sys.exit("Falta el paquete: py -m pip install pypdf")
    pages = []
    for i, page in enumerate(PdfReader(str(path)).pages, 1):
        txt = (page.extract_text() or "").replace("-\n", "")
        pages.append((i, re.sub(r"[ \t]+", " ", txt)))
    chars = sum(len(t.strip()) for _, t in pages)
    print(f"  · {path.name}: {len(pages)} páginas, {chars:,} caracteres de texto")
    if chars < 30 * max(1, len(pages)):
        print(f"  ! {path.name}: casi no tiene texto extraíble (¿PDF escaneado?). No se podrá buscar en él.")
    return pages


def make_chunks(pages, size=1200):
    chunks = []
    for pno, text in pages:
        cur = ""
        for s in re.split(r"(?<=[.!?:])\s+|\n{2,}", text):
            s = s.strip()
            if not s:
                continue
            if cur and len(cur) + len(s) > size:
                chunks.append((pno, cur))
                cur = cur.split(". ")[-1] + " " + s       # pequeño solapamiento
            else:
                cur = (cur + " " + s).strip()
        if cur:
            chunks.append((pno, cur))
    return chunks


class BM25:
    def __init__(self, docs, k1=1.5, b=0.75):
        self.toks = [Counter(tokens(d)) for d in docs]
        self.lens = [sum(t.values()) for t in self.toks]
        self.avg = (sum(self.lens) / len(self.lens)) if self.lens else 1
        df = Counter(w for t in self.toks for w in t)
        n = len(docs)
        self.idf = {w: math.log(1 + (n - f + 0.5) / (f + 0.5)) for w, f in df.items()}
        self.k1, self.b = k1, b

    def scores(self, query_terms):
        q = [w for term in query_terms for w in tokens(term)]
        out = []
        for i, tf in enumerate(self.toks):
            s = 0.0
            for w in q:
                if w in tf:
                    f = tf[w]
                    s += self.idf[w] * f * (self.k1 + 1) / (f + self.k1 * (1 - self.b + self.b * self.lens[i] / self.avg))
            out.append(s)
        return out


def candidates(ranked, sc, chunks, names, k, per_manual, max_chars=60000):
    """Mejores fragmentos globales + mejores de cada manual + sus VECINOS (anterior y siguiente),
    para que un caso o excepción que quedó en el fragmento contiguo no se pierda por el corte."""
    base = [i for i in ranked[:k] if sc[i] > 0]
    for n in names:
        base += [i for i in ranked if chunks[i][0] == n and sc[i] > 0][:per_manual]
    seen, out, total = set(), [], 0
    for i in sorted(set(base), key=lambda i: sc[i], reverse=True):
        for j in (i - 1, i, i + 1):
            if 0 <= j < len(chunks) and j not in seen and chunks[j][0] == chunks[i][0]:
                if total + len(chunks[j][2]) > max_chars:
                    break
                seen.add(j)
                out.append(j)
                total += len(chunks[j][2])
    return sorted(out)       # en orden del documento: se lee como texto continuo


def verify_quote(fragment: str, pages, page_hint):
    """¿El fragmento aparece de verdad en el PDF? Protege contra citas inventadas o parafraseadas."""
    f = norm(fragment)
    if not f:
        return "N/A"
    near = [t for p, t in pages if page_hint is None or abs(p - page_hint) <= 1] or [t for _, t in pages]
    for scope in (near, [t for _, t in pages]):
        full = norm(" ".join(scope))
        if f in full:
            return "SI"
    full = norm(" ".join(near))
    sm = difflib.SequenceMatcher(None, f, full, autojunk=False)
    blk = sm.find_longest_match(0, len(f), 0, len(full))
    win = full[max(0, blk.b - blk.a): blk.b - blk.a + len(f) + 20]
    ratio = difflib.SequenceMatcher(None, f, win, autojunk=False).ratio()
    return "APROX" if ratio >= 0.85 else "NO"


# =========================================================================== Excel

def open_sheet(wb, name, cols):
    if name not in wb.sheetnames:
        sys.exit(f"El Excel debe tener una hoja llamada '{name}'. Crea uno nuevo con: py eval_rag.py plantilla")
    ws = wb[name]
    header = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    col = {h: i + 1 for i, h in enumerate(header) if h}
    for name_ in cols:                          # agrega columnas que falten al final
        if name_ not in col:
            col[name_] = ws.max_column + 1
            ws.cell(row=1, column=col[name_], value=name_).font = Font(bold=True)
    rows = []
    for r in range(2, ws.max_row + 1):
        val = {n: ws.cell(row=r, column=c).value for n, c in col.items()}
        val = {k: ("" if v is None else str(v).strip()) for k, v in val.items()}
        if not any(val.values()):
            continue
        val["_row"] = r
        rows.append(val)
    return ws, col, rows


def read_book(path: Path):
    if not path.exists():
        sys.exit(f"No encuentro {path}. Créalo con: py eval_rag.py plantilla")
    wb = load_workbook(path)
    _, _, fuentes = open_sheet(wb, SH_FUENTES, COLS_FUENTES)
    ws, col, preg = open_sheet(wb, SH_PREG, COLS_PREG)
    fuentes = [f for f in fuentes if f.get("archivo")]
    preg = [p for p in preg if p.get("question")]
    for i, p in enumerate(preg, 1):
        p["id"] = p.get("id") or f"q{i}"
    return wb, ws, col, fuentes, preg


def save_wb(wb, path: Path):
    try:
        wb.save(path)
    except PermissionError:
        sys.exit(f"No pude guardar {path}: ciérralo en Excel y vuelve a correr el comando.")


def is_validated(v: str) -> bool:
    return norm(v) in {"si", "s", "yes", "x", "ok"}


def style_header(ws):
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
    ws.freeze_panes = "A2"


def wrap(ws):
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")


# =========================================================================== comandos

def cmd_plantilla(args):
    path = Path(args.excel)
    if path.exists() and not args.force:
        sys.exit(f"{path} ya existe. Usa --force para sobrescribirlo.")
    wb = Workbook()

    ws = wb.active
    ws.title = SH_FUENTES
    ws.append(COLS_FUENTES)
    ws.append(["reglamento.pdf", "[Título del manual 1]", "[Autor / institución]", "Reemplaza con tus datos"])
    ws.append(["tecnica.pdf", "[Título del manual 2]", "[Autor / institución]", ""])
    ws.append(["entrenamiento.pdf", "[Título del manual 3]", "[Autor / institución]", ""])
    for letter, w in zip("ABCD", (24, 50, 35, 40)):
        ws.column_dimensions[letter].width = w
    style_header(ws)

    ws = wb.create_sheet(SH_PREG)
    ws.append(COLS_PREG)
    ej = [
        ["q1", "¿A qué altura debe estar la red en el centro de la cancha?",
         "Según el documento, la red debe medir 0,914 m (3 pies) en el centro, sujeta por una faja blanca.",
         "", "", "", "", "reglamento.pdf",
         "[EJEMPLO] La red estará suspendida por una cuerda o cable metálico. En el centro, su altura será de 0,914 m (3 pies), donde se mantendrá tensa por medio de una faja de color blanco.",
         "SI"],
        ["q2", "¿Cuáles son las medidas de la cancha para individuales?",
         "La cancha de individuales mide 23,77 m de largo por 8,23 m de ancho. Además, la superficie más rápida es el césped, por eso Wimbledon es el torneo más rápido.",
         "", "", "", "", "reglamento.pdf",
         "[EJEMPLO] La cancha será un rectángulo de 23,77 m de largo y, para partidos de individuales, 8,23 m de ancho. Para dobles, el ancho será de 10,97 m.",
         "SI"],
        ["q3", "¿Qué empuñadura se recomienda para el saque de un principiante?",
         "Se recomienda la empuñadura continental para el saque, porque permite pronar el antebrazo y generar efecto.",
         "", "", "", "", "tecnica.pdf",
         "[EJEMPLO] Para el servicio se aconseja la empuñadura continental, ya que facilita la pronación del antebrazo y la generación de efectos. Los principiantes pueden comenzar con una empuñadura más cerrada y migrar progresivamente.",
         "SI"],
        ["q4", "¿Cuántas horas semanales se recomiendan para un adulto principiante?",
         "Se recomienda entrenar 3 sesiones de 90 minutos por semana, es decir, 4,5 horas semanales.",
         "", "", "", "", NO_FUENTE, "", "SI"],
    ]
    for row in ej:
        ws.append(row)
    for i, w in enumerate([6, 40, 50, 20, 8, 50, 14, 20, 50, 10], 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    style_header(ws)
    wrap(ws)

    ins = wb.create_sheet("Instrucciones")
    for line in [
        "CÓMO USAR ESTE ARCHIVO",
        "",
        "Hoja Fuentes: una fila por PDF que sirve de fuente de verdad. 'archivo' debe ser el nombre exacto del PDF,",
        "   que debe estar en la misma carpeta que este Excel (o en la carpeta que indiques con --pdfs).",
        "   Título y autor no los usa el script; te sirven para tu informe.",
        "",
        "Hoja Preguntas: llena solo id, question (lo que le preguntaste a ChatGPT) y answer (lo que respondió, sin corregir).",
        "   Las 4 filas de ejemplo ya están validadas para que pruebes 'evaluar'. Bórralas al cargar tus datos.",
        "",
        "1. Con el Excel CERRADO:  py eval_rag.py sugerir",
        "   El script busca en TODOS los PDF y llena fuente_sugerida, pagina, fragmento_sugerido y cita_verificada.",
        "   También precarga 'fuente' y 'documents' con la sugerencia.",
        "2. Revisa cada fila abriendo el PDF en esa página:",
        "   - Si es correcto: escribe SI en 'validado'.",
        "   - Si no: corrige 'fuente' y 'documents' con el texto correcto y escribe SI.",
        "   - Si la respuesta no está en ningún manual: 'documents' VACÍO, 'fuente' = (ninguna), y SI.",
        "   - cita_verificada: SI = el texto aparece tal cual en el PDF; APROX = casi igual, revisar;",
        "     NO = no aparece así en el PDF (posible cita inventada), revisar con cuidado.",
        "3. py eval_rag.py evaluar   -> solo evalúa filas con SI. Resultado en resultados.xlsx.",
        "",
        "La columna 'documents' es la fuente de verdad que TÚ validas. El juez compara la respuesta de ChatGPT contra ella.",
    ]:
        ins.append([line])
    ins["A1"].font = Font(bold=True, size=13)
    ins.column_dimensions["A"].width = 125

    wb.active = 1
    save_wb(wb, path)
    print(f"Creado {path}")


def cmd_sugerir(args):
    path = Path(args.excel)
    pdf_dir = Path(args.pdfs) if args.pdfs else path.parent
    wb, ws, col, fuentes, preg = read_book(path)
    if not fuentes:
        sys.exit("La hoja Fuentes está vacía: agrega el nombre de tus PDF en la columna 'archivo'.")

    # Indexa todos los PDF en un solo corpus
    print("Leyendo fuentes:")
    pages_by, chunks = {}, []          # chunks: (archivo, pagina, texto)
    for f in fuentes:
        name = f["archivo"] if f["archivo"].lower().endswith(".pdf") else f["archivo"] + ".pdf"
        pdf = pdf_dir / name
        if not pdf.exists():
            sys.exit(f"No encuentro {pdf}. Revisa el nombre en la hoja Fuentes o usa --pdfs <carpeta>.")
        pages_by[name] = load_pdf(pdf)
        chunks += [(name, p, t) for p, t in make_chunks(pages_by[name])]
    if not chunks:
        sys.exit("No se pudo extraer texto de ningún PDF.")
    bm25 = BM25([t for _, _, t in chunks])
    names = list(pages_by)

    if args.ids:
        wanted = {x.strip() for x in args.ids.split(",") if x.strip()}
        todo = [p for p in preg if p["id"] in wanted]
    else:
        todo = [p for p in preg if not is_validated(p.get("validado", ""))
                and (args.todas or not p.get("documents"))]
    if not todo:
        print("No hay filas pendientes (todas validadas o con 'documents' lleno). Usa --todas para re-sugerir.")
        return
    llm = Claude(args.model, args.dry_run)
    print(f"\nBuscando respuesta para {len(todo)} preguntas en {len(names)} PDF "
          f"({'DRY-RUN' if args.dry_run else args.model}) ...")

    for p in todo:
        # 1) términos de búsqueda (español + inglés)
        terms = [p["question"]]
        if not args.dry_run:
            obj = parse_json(llm.call(KEYWORDS_PROMPT.format(question=p["question"]), max_tokens=1000)) or {}
            terms += [str(t) for t in obj.get("terminos", [])]

        # 2) candidatos + 3) Claude elige el pasaje. Si no encuentra nada, segunda vuelta más amplia.
        sc = bm25.scores(terms)
        ranked = sorted(range(len(chunks)), key=lambda i: sc[i], reverse=True)
        pasajes, comment, vuelta = [], "", 0
        for vuelta, (k, pm) in enumerate([(args.k, args.por_manual), (args.k * 3, args.por_manual * 3)], 1):
            pick = candidates(ranked, sc, chunks, names, k, pm)
            if not pick:
                comment = "Ningún fragmento de los manuales coincide con los términos de búsqueda."
                break
            if args.dry_run:
                a, pg, txt = chunks[max(pick, key=lambda i: sc[i])]
                pasajes = [{"archivo": a, "pagina": pg, "fragmento": re.split(r"(?<=[.!?])\s+", txt)[0]}]
                comment = "dry-run: primera oración del mejor fragmento"
                break
            listing = "\n\n".join(f"[{chunks[i][0]} | p. {chunks[i][1]}] {chunks[i][2]}" for i in pick)
            obj = parse_json(llm.call(LOCATE_PROMPT.format(question=p["question"], chunks=listing),
                                      max_tokens=3000)) or {}
            pasajes, comment = obj.get("pasajes") or [], str(obj.get("comentario", ""))
            if args.debug:
                with open("sugerir_debug.txt", "a", encoding="utf-8") as dbg:
                    dbg.write(f"\n{'=' * 100}\n{p['id']} | vuelta {vuelta} | {p['question']}\n"
                              f"TÉRMINOS: {terms[1:]}\n"
                              f"RESULTADO: {len(pasajes)} pasaje(s) | {comment}\n{'-' * 100}\n"
                              f"FRAGMENTOS QUE VIO CLAUDE ({len(pick)}):\n{listing}\n")
            if pasajes:
                break

        # 4) verificación de cada cita contra su PDF
        found = []
        for ps in pasajes[:3]:
            a = str(ps.get("archivo", "")).strip()
            a = next((n for n in names if norm(n) == norm(a) or norm(n).startswith(norm(a))), a)
            try:
                pg = int(ps.get("pagina"))
            except (TypeError, ValueError):
                pg = None
            frag = str(ps.get("fragmento", "")).strip()
            if not frag:
                continue
            check = verify_quote(frag, pages_by[a], pg) if a in pages_by else "NO (archivo desconocido)"
            found.append((a, pg, frag, check))

        r = p["_row"]
        if found:
            fuente = "; ".join(sorted({a for a, _, _, _ in found}))
            paginas = ", ".join(str(pg) for _, pg, _, _ in found if pg)
            frag_txt = "\n\n".join(f"[{a} p.{pg}] {fr}" if len(found) > 1 else fr for a, pg, fr, _ in found)
            checks = ", ".join(c for _, _, _, c in found)
        else:
            fuente, paginas, checks = NO_FUENTE, "", "NO ENCONTRADO"
            frag_txt = f"(no encontrado) {comment}"
        ws.cell(row=r, column=col["fuente_sugerida"], value=fuente)
        ws.cell(row=r, column=col["pagina"], value=paginas)
        ws.cell(row=r, column=col["fragmento_sugerido"], value=frag_txt)
        ws.cell(row=r, column=col["cita_verificada"], value=checks)
        if not p.get("documents") or (args.ids and not is_validated(p.get("validado", ""))):  # precarga; TÚ validas
            ws.cell(row=r, column=col["fuente"], value=fuente)
            ws.cell(row=r, column=col["documents"], value=frag_txt if found else "")
        for c in ("fuente_sugerida", "pagina", "fragmento_sugerido", "cita_verificada", "fuente", "documents"):
            ws.cell(row=r, column=col[c]).alignment = Alignment(wrap_text=True, vertical="top")
        print(f"  {p['id']:>6}  {fuente:<28} p.{paginas or '-':<8} cita: {checks}")

    save_wb(wb, path)
    print(f"\nListo. Abre {path}, revisa cada fila contra el PDF y marca SI en 'validado'.")


def cmd_evaluar(args):
    path = Path(args.excel)
    _, _, _, fuentes, preg = read_book(path)
    metrics = [m.strip() for m in args.metrics.split(",") if m.strip()]
    if set(metrics) - set(METRICS_ALL):
        sys.exit(f"Métricas válidas: {', '.join(METRICS_ALL)}")
    ok = [p for p in preg if is_validated(p.get("validado", ""))]
    skipped = len(preg) - len(ok)
    if not ok:
        sys.exit("No hay filas con SI en 'validado'. Valida los fragmentos antes de evaluar.")
    for p in ok:
        if not p.get("answer"):
            sys.exit(f"La fila {p['id']} no tiene 'answer'.")
    llm = Claude(args.model, args.dry_run)
    print(f"Evaluando {len(ok)} preguntas validadas ({'DRY-RUN' if args.dry_run else args.model})"
          + (f"; se omiten {skipped} sin validar" if skipped else "") + " ...")

    results = []
    for p in ok:
        docs = p.get("documents", "") or "(no se proporcionaron documentos)"
        fuente = p.get("fuente") or (NO_FUENTE if not p.get("documents") else "(sin indicar)")
        out = {"id": p["id"], "fuente": fuente, "question": p["question"],
               "answer": p["answer"], "documents": p.get("documents", "")}
        reference = ""
        if "relevance" in metrics:
            reference = "Respuesta de referencia simulada (dry-run)."
            if not args.dry_run:
                for _ in range(3):                      # reintenta si vuelve vacía
                    reference = llm.call(REFERENCE_PROMPT.format(documents=docs, question=p["question"]),
                                         max_tokens=2000).strip()
                    if reference:
                        break
                reference = reference or "(el modelo no devolvió referencia)"
            out["reference_answer"] = reference
        for m in metrics:
            if args.dry_run:
                score, reason = random.randint(1, 5), "Calificación simulada (dry-run)."
            else:
                prompt = RUBRICS[m].format(question=p["question"], answer=p["answer"],
                                           documents=docs, reference=reference)
                score, reason = None, ""
                for _ in range(3):                      # reintenta si la respuesta no se pudo leer
                    score, reason = parse_score(llm.call(prompt, system=SYSTEM_JUDGE, max_tokens=2000))
                    if score is not None:
                        break
                if score is None:
                    print(f"  ! {p['id']} {m}: no se pudo leer la calificación del juez")
            out[m], out[f"{m}_reason"] = score, reason
        results.append(out)
        print(f"  {p['id']:>6}  " + "  ".join(f"{m}={out[m]}" for m in metrics))

    write_results(Path(args.out), results, metrics, fuentes)
    print(f"\nListo: {args.out}")


def write_results(path: Path, results, metrics, fuentes):
    wb = Workbook()
    ws = wb.active
    ws.title = "Resultados"
    fields = list(results[0].keys())
    ws.append(fields)
    for r in results:
        ws.append([r.get(f) for f in fields])
    red, amber, green = (PatternFill("solid", fgColor=c) for c in ("F8CBAD", "FFE699", "C6EFCE"))
    style_header(ws)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if fields[c.column - 1] in metrics and isinstance(c.value, int):
                c.fill = red if c.value <= 2 else amber if c.value == 3 else green
                c.alignment = Alignment(horizontal="center", vertical="top")
    for i, f in enumerate(fields, 1):
        w = 10 if f in metrics or f == "id" else 22 if f == "fuente" else 45
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    ws.freeze_panes = "B2"

    rs = wb.create_sheet("Resumen")
    rs.append(["Fuente", "n"] + metrics)
    groups = {}
    for r in results:
        groups.setdefault(r["fuente"], []).append(r)

    def avg(items, m):
        v = [x[m] for x in items if isinstance(x.get(m), int)]
        return round(sum(v) / len(v), 2) if v else None

    for name, items in sorted(groups.items()):
        rs.append([name, len(items)] + [avg(items, m) for m in metrics])
    rs.append(["TOTAL", len(results)] + [avg(results, m) for m in metrics])
    for c in rs[1]:
        c.font = Font(bold=True)
    for c in rs[rs.max_row]:
        c.font = Font(bold=True)
    rs.column_dimensions["A"].width = 30

    low = [r for r in results if any(isinstance(r.get(m), int) and r[m] <= 2 for m in metrics)]
    if low:
        rs.append([])
        rs.append(["Respuestas con alguna métrica ≤ 2 (revisar a mano)"])
        rs.cell(row=rs.max_row, column=1).font = Font(bold=True)
        for r in low:
            for m in metrics:
                if isinstance(r.get(m), int) and r[m] <= 2:
                    rs.append([r["id"], m, r[m], r[f"{m}_reason"]])

    fs = wb.create_sheet("Fuentes")
    fs.append(COLS_FUENTES)
    for f in fuentes:
        fs.append([f.get(c, "") for c in COLS_FUENTES])
    style_header(fs)
    for letter, w in zip("ABCD", (24, 50, 35, 40)):
        fs.column_dimensions[letter].width = w
    save_wb(wb, path)


# =========================================================================== main

def main():
    ap = argparse.ArgumentParser(description="Evaluador RAG con Claude como juez y validación humana")
    sub = ap.add_subparsers(dest="cmd", required=True)

    p = sub.add_parser("plantilla", help="Crea el Excel de ejemplo")
    p.add_argument("--excel", default="tenis_eval.xlsx")
    p.add_argument("--force", action="store_true")

    p = sub.add_parser("sugerir", help="Busca en todos los PDF el fragmento que responde cada pregunta")
    p.add_argument("--excel", default="tenis_eval.xlsx")
    p.add_argument("--pdfs", help="Carpeta de los PDF (por defecto, la del Excel)")
    p.add_argument("--k", type=int, default=12, help="Mejores fragmentos globales que ve Claude")
    p.add_argument("--por-manual", type=int, default=4, help="Mejores fragmentos garantizados por manual")
    p.add_argument("--todas", action="store_true", help="Re-sugerir también filas con 'documents' lleno")
    p.add_argument("--ids", help="Solo estas preguntas, separadas por coma (ej. q1,q7)")
    p.add_argument("--debug", action="store_true", help="Guarda en sugerir_debug.txt lo que vio Claude")
    p.add_argument("--model", default=DEFAULT_MODEL)
    p.add_argument("--dry-run", action="store_true")

    p = sub.add_parser("evaluar", help="Evalúa las filas validadas")
    p.add_argument("--excel", default="tenis_eval.xlsx")
    p.add_argument("--metrics", default=",".join(METRICS_ALL))
    p.add_argument("--out", default="resultados.xlsx")
    p.add_argument("--model", default=DEFAULT_MODEL)
    p.add_argument("--dry-run", action="store_true")

    args = ap.parse_args()
    {"plantilla": cmd_plantilla, "sugerir": cmd_sugerir, "evaluar": cmd_evaluar}[args.cmd](args)


if __name__ == "__main__":
    main()
